#!/usr/bin/env python3
"""Generate the UBCC delivery deck; run in ubcc-dev:ubuntu20.04.

Dependencies: python-pptx, Pillow; --render additionally needs soffice and
PyMuPDF. Example: python3 scripts/generate_delivery_overview_ppt.py
  --output-dir /output --render
All reported measurements are transcribed from the three design deliverables.
"""
import argparse
import json
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / 'docs/design/figures'
NAVY, BLUE, ORANGE = '142D4E', '285E8E', 'DC8738'
GRAY, LIGHT, WHITE = '566579', 'EDF2F7', 'FFFFFF'
FONT = 'Microsoft YaHei'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
titles = []


def rect(s, x, y, w, h, color):
    sh = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, value, size=22, color=NAVY, bold=False):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.03)
    tf.margin_top = tf.margin_bottom = Inches(.03)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size, p.font.bold = FONT, Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(0)
        for run in p.runs:
            rp = run._r.get_or_add_rPr()
            ea = OxmlElement('a:ea')
            ea.set('typeface', FONT)
            rp.append(ea)
    return box


def slide(title, section, source):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    titles.append(title)
    rect(s, 0, 0, 16, .12, NAVY)
    text(s, .65, .35, 14, .35, section, 13, BLUE, True)
    text(s, .65, .95, 14.7, .7, title, 32, NAVY, True)
    rect(s, .68, 1.84, .65, .055, ORANGE)
    text(s, .68, 8.48, 14, .25, source, 11, GRAY)
    text(s, 14.7, 8.43, .7, .35, '%02d / 16' % len(titles), 11, GRAY)
    s.notes_slide.notes_text_frame.text = source
    return s


def bullets(s, items, x=.8, y=2.2, w=14.4, step=1.04, size=23):
    for i, item in enumerate(items):
        rect(s, x, y+i*step+.15, .08, .08, ORANGE)
        text(s, x+.25, y+i*step, w-.25, step-.1, item, size)


def table(s, headers, rows, widths, x=.7, y=2.2, row=.64, size=19):
    for i, cells in enumerate([headers] + rows):
        xx = x
        for cell, width in zip(cells, widths):
            rect(s, xx, y+i*row, width-.025, row-.025,
                 NAVY if i == 0 else (LIGHT if i % 2 else 'F7F9FC'))
            text(s, xx+.12, y+i*row+.09, width-.23, row-.12,
                 str(cell), size, WHITE if i == 0 else NAVY, i == 0)
            xx += width


def pic(s, name, x, y, w, h):
    p = FIG / name
    with Image.open(p) as im:
        iw, ih = im.size
    scale = min(w/iw, h/ih)
    ww, hh = iw*scale, ih*scale
    s.shapes.add_picture(str(p), Inches(x+(w-ww)/2), Inches(y+(h-hh)/2),
                         width=Inches(ww), height=Inches(hh))


def card(s, x, y, w, label, value, detail):
    rect(s, x, y, w, 1.65, LIGHT)
    text(s, x+.2, y+.15, w-.4, .35, label, 16, GRAY)
    text(s, x+.2, y+.55, w-.4, .65, value, 32, BLUE, True)
    text(s, x+.2, y+1.24, w-.4, .3, detail, 14, GRAY)


def build():
    s = slide('UBCC 跨节点缓存一致性方案', 'UBCC  /  三项交付件成果汇报', 'UBCC · 交付总览')
    rect(s, 0, 2.1, 16, 4.5, NAVY)
    text(s, .95, 2.75, 14, 1, '独立全局目录，连接跨节点一致性', 38, WHITE, True)
    text(s, .95, 4.25, 14, .65, '协议体系结构 · 形式化验证与可靠性 · 性能验收与集成接口', 24, WHITE)
    text(s, 1, 7.15, 14, .55, '分层元数据容量扩展   /   精确权限仲裁   /   可验证的集成路径', 23, BLUE)

    s = slide('交付总览与三项指标结论', '总览', '来源：D1 §1；D2 §1、§5；D3 §1、§5')
    text(s, .8, 2.1, 14.4, .85, 'D1 定义架构与协议语义    D2 验证正确性与故障恢复    D3 核验性能并定义集成接口', 21)
    table(s, ['指标', '结果', '验收门槛 / 结论'], [
        ['1 · 等效追踪容量', '1.515×', '≥1.500×，达标'],
        ['1 · 目录溢出附加时延', '10.535 ns', '<25 ns，达标'],
        ['2 · 适用场景等权平均降幅', '64.759%', '≥10%，达标'],
        ['3 · UBCC 相对 HA-VI', '核心组 20.8%–27.1%\n代表组 8.7%–25.0%', '五拓扑 × 三压力\n15 个坐标的两个组全部满足']], [4.5, 4.65, 5.45], y=3.05, row=.82, size=20)
    text(s, .8, 7.35, 14.4, .8, '正确性门禁全部通过：72/72（指标1/2） · 6/6（独立时延） · 240/240（配对）\n6/6（重型回归） · 52/52（故障资格）', 19, BLUE, True)

    s = slide('独立 Outer 层，分离全局与节点内职责', 'D1  /  方案概述', '来源：D1 §1–§3；图：ubcc-system-architecture.png')
    pic(s, 'ubcc-system-architecture.png', .65, 2.15, 10.1, 5.95)
    bullets(s, ['节点内 gem5 CHI 保持本地一致性职责', 'EP 边界转换请求、数据与完成', 'Home UBCC 控制器统一全局权限仲裁', 'ResidentDir 与 H64 Backstore 独立管理元数据容量'], x=11, y=2.35, w=4.2, step=1.37, size=21)

    s = slide('分层目录：热点驻留，冷元数据扩容', 'D1  /  ResidentDir + Bloom + H64 Backstore', '来源：D1 §3.2；图：ubcc-metadata-fanout-scaling.png')
    pic(s, 'ubcc-metadata-fanout-scaling.png', .65, 2.05, 10, 5.95)
    bullets(s, ['ResidentDir：片上 SRAM\ntag / state / epoch / sharers', 'Bloom：分组过滤；三条查询路径为 ResDir Hit、ResDir Miss+BF Hit、BF Miss', 'H64 Backstore：metadata DRAM；64 B bucket 开放寻址哈希；仅存目录元数据', '固定 512 KiB 片上预算\n有效覆盖按地址去重计算'], x=10.95, y=2.2, w=4.4, step=1.47, size=19)

    s = slide('全局 MESI 与两阶段提交', 'D1  /  全局一致性语义', '来源：D1 §3–§4；D3 §2.6、附录 F')
    table(s, ['G_I', 'G_S', 'G_E', 'G_M'], [['无有效副本', '共享只读', '干净独占', '修改独占']], [3.65]*4, y=2.2, row=.6, size=22)
    bullets(s, ['owner 由 one-hot sharer 推导；精确选择 Recall / Invalidate 目标', 'Grant 阶段保留权限；committed 与 intended 分离', '匹配 Clear 后提交目标状态，并退役对应 outstanding', 'epoch / reqId 单调与身份校验：旧消息不推进新事务', '单向完成语义：数据与权限可观察即根操作完成，不等待同步 ClearResp'], y=3.7, step=.84, size=22)

    s = slide('三条关键路径，共用一个提交点', 'D1  /  远程读 · 所有权迁移 · 共享转写者', '来源：D1 §5；图：ubcc-protocol-paths.png')
    # Lay out the three complete panels side by side, retaining the source
    # aspect ratio within each crop (all source pixels remain represented).
    path = FIG / 'ubcc-protocol-paths.png'
    with Image.open(path) as im:
        iw, ih = im.size
    for i, (label, desc) in enumerate([
            ('远程读', '定位权威数据源\n返回数据与共享授权'),
            ('所有权迁移', '旧 owner 释放权限与最新数据\n新 owner 获得单一写权限'),
            ('共享转写者', '精确失效实际 sharer\nAck 集合收敛后授权')]):
        x = .75 + i*4.95
        text(s, x, 2.25, 4.65, .5, label, 25, BLUE, True)
        text(s, x, 2.93, 4.65, .9, desc, 21)
        image = s.shapes.add_picture(str(path), Inches(x), Inches(4.2),
                                    width=Inches(4.65), height=Inches(4.65*ih/iw/3))
        image.crop_top = i/3
        image.crop_bottom = (2-i)/3
    text(s, .9, 7.5, 14.2, .6, '先完成数据回收与目标确认，再授予权限；requester 本地完成后发送 Clear，Home 校验并提交。', 22, BLUE, True)

    s = slide('方案选择：精确全局 MESI 目录＋边界转换', 'D1  /  全局目录方案与状态集合', '来源：D1 §7.2–§7.6；机制对比，不作为额外性能测量')
    table(s, ['维度', '多节点系统级 CHI HN-F', '全局 MESI 目录＋边界协议转换'], [
        ['存储', '全局目录＋CHI 通道与事务字段', '2-bit 状态＋节点位图；冷容量分层扩展'],
        ['时延', '原生 CHI；可用 DMT/DCT 直接供数', '边界映射；远程供数经 Home；E 优化私有写'],
        ['流量', '取决于过滤精度、snoop 与数据路由', '精确失效降低控制流量；中转增加数据遍历'],
        ['复杂度', '复用标准资产；管理通道、credit 与 ID', '语义集中；EP 负责跨域完成关联'],
        ['状态转换', '依 CHI 版本、目录及 agent 能力', 'I/S/E/M 权限清晰，提交点明确']], [1.45, 6.25, 6.9], row=.66, size=18)
    text(s, .85, 6.45, 14.2, .75, '选择 MESI：E 覆盖常见私有读后写；与 MSI 同为 2-bit；I/S/E/M 覆盖主要权限关系。', 22, BLUE, True)
    text(s, .85, 7.35, 14.2, .8, '按瓶颈引入增强：远程中转→直接转发；持续脏共享→O；干净共享供数→F；\n已有系统级 CHI fabric 或互操作需求→全局 HN-F。', 19)

    s = slide('从形式化模型到故障资格的分层验证', 'D2  /  验证体系与总体结论', '来源：D2 §1–§2、§7；图：ubcc-verification-stack.png')
    pic(s, 'ubcc-verification-stack.png', .65, 2.2, 7.2, 5.9)
    table(s, ['验证内容', '结果'], [
        ['目录核心状态与提交顺序', '通过'], ['Recall / Invalidate / Upgrade / Clear 活性', '通过'],
        ['EP-RNF snoop 仲裁', '通过'], ['多 PA / 多 Socket 隔离', '通过'],
        ['O3 处理器同步语义', '通过'], ['丢包、重复、乱序故障资格', '52/52']], [5.8, 1.4], x=8.1, y=2.4, row=.76, size=18)
    text(s, 8.2, 7.9, 7, .35, '拓扑覆盖：3N1S、3N2S、8N2S、16N1S', 17, BLUE)

    s = slide('形式化结果：安全性与活性分别核验', 'D2  /  TLA+ 与定向机制验证', '来源：D2 §3；focused 模型活性由定向可执行验证支持')
    table(s, ['模型组', 'Safety', 'Liveness / 验证方式'], [
        ['UBCC 目录核心', '通过', '通过'], ['传输故障模型', '通过', '通过'],
        ['多 PA / 多 Socket 隔离', '通过', '通过'],
        ['EP-RNF snoop 仲裁', '通过 · 328 个状态', '定向端到端验证通过'],
        ['committed waiter 精确退役', '通过 · 274,593 个状态', '退役与升级重放定向验证通过']], [5.6, 4.1, 4.9], row=.77, size=21)
    bullets(s, ['Safety：单一写 owner、epoch 不回退、重复 Ack 不重复推进、匹配 Clear 才提交', 'Liveness：公平调度、消息最终可达条件下，请求推进、确认收敛与 waiter 重放'], y=7.02, step=.59, size=18)

    s = slide('可靠性机制覆盖 52 项故障资格', 'D2  /  可恢复传输故障', '来源：D2 §4–§5；同一测试可覆盖多个故障动作类别')
    text(s, .8, 2.15, 14.4, .8, '两阶段提交 · epoch/reqId 校验 · Ack 位图 · tombstone 幂等 · waiter 去重与重放', 22, BLUE, True)
    table(s, ['资格组', '覆盖', '通过'], [
        ['Q1', '基础消息故障', '20/20'], ['Q2', '连续丢失：首 2 次 / 首 3 次', '8/8'],
        ['Q3', '请求与响应双故障组合', '4/4'], ['Q4', '32 PA、突发、部分 Ack、多来源与并发边界', '8/8'],
        ['Q5', '3N1S、3N2S、8N2S、16N1S', '12/12']], [1.7, 10.4, 2.5], y=3.1, row=.66, size=20)
    text(s, .8, 7.2, 14.4, .85, '消息路径：Clear / Upgrade / Invalidate / Recall    总计：52/52\n故障动作覆盖：丢失 37 · 延迟 13 · 重复 6 · 乱序 6（类别可重叠）', 21, BLUE, True)

    s = slide('指标 1：容量扩展与附加时延均达标', 'D3  /  等效追踪容量与目录溢出成本', '来源：D3 §2.3、§3；扩展观测按 TC 合并，容量 GM / 附加时延 AM')
    card(s, .7, 2.1, 7.15, '固定 512 KiB 预算 · 容量', '1.515×', '99,293 / 65,536；门槛 ≥1.500×')
    card(s, 8.1, 2.1, 7.15, '目录溢出附加时延 · 三轮配对', '10.535 ns', '同协议策略：512 KiB 溢出路径 − 2 MiB 无溢出参考；门槛 <25 ns')
    table(s, ['扩展应用', '容量 GM', '附加时延 AM（ns）'], [
        ['OLTP 缓冲池（TC142）', '1.843×', '14.448'], ['索引页访问（TC143）', '1.618×', '36.967'],
        ['WAL（TC144）', '1.848×', '17.722'], ['FaaS 热调用（TC145）', '1.807×', '25.249'],
        ['图 frontier（TC146）', '1.849×', '17.142'], ['Feature store（TC147）', '1.812×', '23.233'],
        ['合并均值', '1.794×', '22.460']], [6.7, 3.3, 4.6], y=4.0, row=.45, size=17)
    text(s, .85, 7.8, 14.2, .45, '五拓扑 × 175% / 200% 目录压力：60 坐标、180 次运行；六个应用合并容量均 ≥1.5×。', 18)

    s = slide('指标 2：适用场景等权平均降幅 64.759%', 'D3  /  应用端到端时延与服务阶段扩展', '来源：D3 §2.4、§4.1–§4.6；扩展服务阶段不含逐批压力写入与屏障')
    text(s, .8, 2.15, 14.4, .65, '合同值 64.759% ≥10%：optimized 对比 naive，六个适用场景按降幅等权。', 24, BLUE, True)
    table(s, ['合同计分场景', '降幅'], [
        ['已有 sharer 重访', '98.305%'], ['已有 owner 写入', '96.667%'],
        ['新 requester 读取', '25.000%'], ['脏 owner 交接', '−13.333%'],
        ['混合批处理', '97.302%'], ['Catalog 批处理', '84.615%']], [4.8, 2.2], y=3.15, row=.56, size=19)
    table(s, ['扩展服务阶段（按 TC 合并）', '降幅'], [
        ['OLTP 缓冲池（TC142）', '44.330%'], ['索引页（TC143）', '48.760%'],
        ['WAL（TC144）', '37.470%'], ['FaaS（TC145）', '42.290%'],
        ['图 frontier（TC146）', '46.930%'], ['Feature store（TC147）', '50.520%']], [5.25, 2.05], x=8, y=3.15, row=.56, size=18)
    text(s, .8, 7.35, 7, .8, '适用门槛：naive ≥500 ns；\n119.209 ns 的中性控制项不计分。', 19)
    text(s, 8.1, 7.35, 7.1, .8, '扩展 37.5%–50.5%，全部 ≥10%；\n三种单 Socket 拓扑、两档目录压力等权。', 19, BLUE, True)

    s = slide('指标 3：15 个坐标的两个场景组全部满足', 'D3  /  UBCC 与 HA-VI 配对比较', '来源：D3 §5；表内为组均值时延降幅；三档压力为不同条件')
    text(s, .8, 2.12, 14.4, .6, '核心组 20.8%–27.1%   |   代表组 8.7%–25.0%   |   120 对 / 240 次运行', 25, BLUE, True)
    table(s, ['拓扑', '0% L3 压力', '50% L3 压力', '100% L3 压力'], [
        ['3N1S', '22.451% / 8.679%', '24.103% / 8.927%', '27.129% / 8.978%'],
        ['3N2S', '21.247% / 11.685%', '20.814% / 12.570%', '20.911% / 11.769%'],
        ['8N1S', '22.443% / 14.799%', '23.834% / 9.503%', '27.102% / 9.617%'],
        ['8N2S', '21.210% / 16.117%', '20.846% / 14.352%', '24.285% / 14.407%'],
        ['16N1S', '22.722% / 25.011%', '23.900% / 10.881%', '27.070% / 11.072%']], [2.0, 4.2, 4.2, 4.2], y=3.05, row=.66, size=20)
    text(s, .85, 7.2, 14.2, .9, '每格：核心组 / 代表组；各组分别按 3 / 5 个场景等权，判据为 UBCC 组均值严格更低。\n共同条件：O3、256 KiB L3、相同输入与根操作完成边界、单向完成语义。', 20)

    s = slide('集成接口与六步流程', 'D3  /  UBCC · EP-RNF · EP-SNF · EPBackend · UBAdapter', '来源：D3 §7–§8；EPBackend 职责见 D1 §3.7')
    table(s, ['模块 / 接口', '职责与完成关系'], [
        ['UBCC 目录服务', '读 / 写权限 / Clear / Writeback / Evict：授权、提交与目录生命周期'],
        ['EP-RNF', 'ReadShared / ReadUnique / CleanUnique；snoop immediate / stale 仲裁'],
        ['EP-SNF / EPBackend', '节点内服务请求与 CHI 响应；关联 Recall、写回及跨域完成'],
        ['UBAdapter', '发送、接收、回调、稳定重试；保持 epoch / reqId']], [4.1, 10.5], row=.65, size=19)
    text(s, .8, 5.75, 14.4, .85, '公共字段：消息类型 · 源/目标 node 与 socket · PA · epoch/reqId · 权限/状态\n数据有效性与负载 · target mask 与 Ack', 21, BLUE)
    steps = ['确定拓扑\n与地址映射', '装载 gem5\nUBIO 与参考模型', '选择矩阵\n分配并行资源', '按共同身份\n交换事务事件', '检查数据\n完成与模块状态', '按正式口径\n汇总结论']
    for i, st in enumerate(steps):
        x = .75 + i*2.45
        rect(s, x, 7.05, 2.3, 1.0, LIGHT)
        text(s, x+.1, 7.14, 2.12, .85, str(i+1)+'  '+st, 16, NAVY, True)

    s = slide('三册交付内容与适用范围', '交付内容', '来源：D1 §1、§8；D2 §4.7、§8；D3 §1.2、§5.2、§8')
    table(s, ['交付件', '主要内容'], [
        ['D1 · 协议体系结构', '架构、组件、分层目录、一致性语义、关键路径与方案选择'],
        ['D2 · 形式化验证与可靠性', 'TLA+、定向机制、端到端、O3、多拓扑及故障资格矩阵'],
        ['D3 · 性能验收与集成接口', '三项指标、计量口径、配对结果、模块接口与六步集成流程']], [4.4, 10.2], row=.87, size=22)
    bullets(s, ['建模层级：逻辑仿真与参考模型，验证协议节点规模与端点能力', '可靠性条件：节点与 Home 持续存活，可恢复传输故障最终解除', '目标硬件与端口级 Switch 微体系结构未建模'], y=6.12, step=.67, size=22)

    s = slide('三项指标达标，验证与集成路径明确', '总结', '来源：D1、D2、D3 交付结论')
    card(s, .75, 2.35, 4.65, '指标 1 · 固定预算的分层目录', '1.515×', '附加时延 10.535 ns')
    card(s, 5.68, 2.35, 4.65, '指标 2 · 适用场景时延', '64.759%', '等权平均降幅，超过 10% 门槛')
    card(s, 10.61, 2.35, 4.65, '指标 3 · 配对组均值', '15 / 15', '两个场景组均低于 HA-VI')
    bullets(s, ['体系结构明确：独立全局 MESI、精确目录、ResidentDir 与 H64 Backstore', '验证体系确立：形式化、定向、端到端与故障资格；52/52 故障项通过', '集成路径完整：模块接口、事务身份、完成条件与六步验证流程'], y=4.75, step=1.0, size=25)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output-dir', type=Path, default=Path('/mnt/data2/cgc/fileserver'))
    ap.add_argument('--render', action='store_true')
    args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    build()
    assert len(prs.slides) == 16
    for s in prs.slides:
        for sh in s.shapes:
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= prs.slide_width + 10
            assert sh.top + sh.height <= prs.slide_height + 10
    target = args.output_dir / 'UBCC_delivery_overview.pptx'
    prs.save(target)
    if args.render:
        subprocess.run(['soffice', '-env:UserInstallation=file:///tmp/ubcc-ppt-profile',
                        '--headless', '--convert-to', 'pdf', '--outdir', str(args.output_dir),
                        str(target)], check=True)
        import fitz
        doc = fitz.open(str(target.with_suffix('.pdf')))
        assert len(doc) == 16
        review = args.output_dir / 'UBCC_delivery_overview_review'
        review.mkdir(exist_ok=True)
        for i, page in enumerate(doc):
            page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4)).save(str(review / ('page-%02d.png' % (i+1))))
            for block in page.get_text('dict')['blocks']:
                if block['type'] == 0:
                    x0, y0, x1, y1 = block['bbox']
                    assert x0 >= 0 and y0 >= 0 and x1 <= page.rect.width and y1 <= page.rect.height
        (review / 'manifest.json').write_text(json.dumps({'pages': len(doc), 'titles': titles}, ensure_ascii=False, indent=2))
    print(str(target))
    print('\n'.join('%02d %s' % (i+1, t) for i, t in enumerate(titles)))


if __name__ == '__main__':
    main()
